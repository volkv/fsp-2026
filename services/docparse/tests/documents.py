"""Builders of synthetic test documents. People, phones and addresses in them are fictional."""

from __future__ import annotations

import io
import struct
import subprocess
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import IO

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from conftest import fixture_bytes

TABLE_DOCX = "table.docx"
DOCX_NAME_IN_BODY = "Образцова Марина Сергеевна"


def rewrite_docx(
    source: bytes,
    replace: dict[str, bytes],
    *,
    extra: dict[str, bytes] | None = None,
    stream_member: tuple[str, Callable[[IO[bytes]], None]] | None = None,
) -> bytes:
    """Copy a DOCX, replacing members; `stream_member` writes one member through a callback (for huge members)."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(source)) as zin,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zout,
    ):
        for info in zin.infolist():
            if stream_member and info.filename == stream_member[0]:
                with zout.open(info.filename, "w", force_zip64=True) as member:
                    stream_member[1](member)
                continue
            zout.writestr(info.filename, replace.get(info.filename, zin.read(info)))
        for name, data in (extra or {}).items():
            zout.writestr(name, data)
    return out.getvalue()


def document_xml(source: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(source)) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def with_doctype(doctype: str, body_replacement: str) -> bytes:
    """table.docx with a document type declaration and the candidate's name replaced by entity references."""
    source = fixture_bytes(TABLE_DOCX)
    declaration, rest = document_xml(source).split("?>", 1)
    xml = declaration + "?>" + doctype + rest.replace(DOCX_NAME_IN_BODY, body_replacement, 1)
    return rewrite_docx(source, {"word/document.xml": xml.encode("utf-8")})


def xxe_docx(url: str) -> bytes:
    doctype = (
        f'<!DOCTYPE w:document [<!ENTITY local SYSTEM "file:///etc/hostname"><!ENTITY remote SYSTEM "{url}/xxe">]>'
    )
    return with_doctype(doctype, "Before [&local;] [&remote;] after")


def entity_bomb_docx() -> bytes:
    levels = "".join(f'<!ENTITY l{i} "{f"&l{i - 1};" * 10}">' for i in range(1, 10))
    return with_doctype(f'<!DOCTYPE w:document [<!ENTITY l0 "lol">{levels}]>', "&l9;")


def utf16_doctype_docx() -> bytes:
    source = fixture_bytes(TABLE_DOCX)
    declaration, rest = document_xml(source).split("?>", 1)
    xml = declaration.replace("UTF-8", "UTF-16") + '?><!DOCTYPE w:document [<!ENTITY x "y">]>' + rest
    return rewrite_docx(source, {"word/document.xml": xml.encode("utf-16")})


def paragraphs_docx(paragraphs_count: int) -> bytes:
    """A DOCX whose body is `paragraphs_count` empty paragraphs: a few bytes each in XML, far more once parsed."""
    source = fixture_bytes(TABLE_DOCX)
    head = document_xml(source).split("<w:body>")[0] + "<w:body>"
    chunk = b"<w:p/>" * 10_000

    def write(member: IO[bytes]) -> None:
        member.write(head.encode("utf-8"))
        for _ in range(paragraphs_count // 10_000):
            member.write(chunk)
        member.write(b"<w:p><w:r><w:t>end</w:t></w:r></w:p></w:body></w:document>")

    return rewrite_docx(source, {}, stream_member=("word/document.xml", write))


def filler_docx(megabytes: int) -> bytes:
    """A DOCX with `megabytes` of repeated text in the body: tiny compressed, large expanded."""
    source = fixture_bytes(TABLE_DOCX)
    head = document_xml(source).split("<w:body>")[0] + "<w:body>"
    paragraph = b"<w:p><w:r><w:t>" + b"A" * 1000 + b"</w:t></w:r></w:p>"

    def write(member: IO[bytes]) -> None:
        member.write(head.encode("utf-8"))
        for _ in range(megabytes * 1024 * 1024 // len(paragraph) + 1):
            member.write(paragraph)
        member.write(b"</w:body></w:document>")

    return rewrite_docx(source, {}, stream_member=("word/document.xml", write))


def lie_about_size(data: bytes, member: str, declared: int) -> bytes:
    """Patch the local header and the central directory so `member` claims `declared` uncompressed bytes."""
    patched = bytearray(data)
    name = member.encode()
    offset = patched.find(b"PK\x03\x04")
    while offset != -1:
        name_len = struct.unpack_from("<H", patched, offset + 26)[0]
        if patched[offset + 30 : offset + 30 + name_len] == name:
            struct.pack_into("<I", patched, offset + 22, declared)
        offset = patched.find(b"PK\x03\x04", offset + 4)
    offset = patched.find(b"PK\x01\x02")
    while offset != -1:
        name_len = struct.unpack_from("<H", patched, offset + 28)[0]
        if patched[offset + 46 : offset + 46 + name_len] == name:
            struct.pack_into("<I", patched, offset + 24, declared)
        offset = patched.find(b"PK\x01\x02", offset + 4)
    return bytes(patched)


def many_entries_zip(entries: int) -> bytes:
    source = fixture_bytes(TABLE_DOCX)
    return rewrite_docx(source, {}, extra={f"customXml/item{i}.bin": b"" for i in range(entries)})


def remote_docx(url: str) -> bytes:
    source = fixture_bytes(TABLE_DOCX)
    with zipfile.ZipFile(io.BytesIO(source)) as archive:
        rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")
    base = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rels = rels.replace(
        "</Relationships>",
        f'<Relationship Id="rIdX1" Type="{base}/image" Target="{url}/img.png" TargetMode="External"/>'
        f'<Relationship Id="rIdX2" Type="{base}/attachedTemplate" Target="{url}/t.dotm" TargetMode="External"/>'
        "</Relationships>",
    )
    return rewrite_docx(source, {"word/_rels/document.xml.rels": rels.encode("utf-8")})


def text_docx(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def text_box_docx(body: str, box: str) -> bytes:
    """A DOCX with a text box written the way Word does: DrawingML content plus a VML fallback copy."""
    source = text_docx([body])
    xml = document_xml(source)
    box_paragraph = f"<w:p><w:r><w:t>{box}</w:t></w:r></w:p>"
    alternate = (
        '<w:r><mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:Choice Requires="wps"><w:drawing>'
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
        '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:txbx>'
        f"<w:txbxContent>{box_paragraph}</w:txbxContent>"
        "</wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
        '<mc:Fallback><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox>'
        f"<w:txbxContent>{box_paragraph}</w:txbxContent>"
        "</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
    )
    xml = xml.replace(f"<w:t>{body}</w:t></w:r>", f"<w:t>{body}</w:t></w:r>{alternate}", 1)
    return rewrite_docx(source, {"word/document.xml": xml.encode("utf-8")})


def blank_png() -> bytes:
    out = io.BytesIO()
    Image.new("L", (800, 1100), color=255).save(out, format="PNG")
    return out.getvalue()


def blank_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


ODT_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
  xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
  xmlns:xlink="http://www.w3.org/1999/xlink" xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
  office:version="1.3">
<office:body><office:text>
<text:h text:outline-level="1">Тестов Иван Петрович</text:h>
<text:p>Инженер по данным, Казань</text:p>
{extra}
</office:text></office:body></office:document-content>"""

ODT_MANIFEST = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">
<manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>
<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
</manifest:manifest>"""


def odt(linked_image_url: str | None = None) -> bytes:
    extra = ""
    if linked_image_url:
        extra = (
            '<text:p><draw:frame svg:width="2cm" svg:height="2cm">'
            f'<draw:image xlink:href="{linked_image_url}" xlink:type="simple"'
            ' xlink:show="embed" xlink:actuate="onLoad"/>'
            "</draw:frame></text:p>"
        )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.text")
        archive.writestr("META-INF/manifest.xml", ODT_MANIFEST)
        archive.writestr("content.xml", ODT_CONTENT.format(extra=extra), compress_type=zipfile.ZIP_DEFLATED)
    return out.getvalue()


def rtf_with_linked_picture(url: str) -> bytes:
    return (
        "{\\rtf1\\ansi{\\fonttbl{\\f0 Arial;}}\\pard Linked picture below\\par"
        f'{{\\field{{\\*\\fldinst INCLUDEPICTURE "{url}/pixel.png" \\\\d}}{{\\fldrslt }}}}\\par}}'
    ).encode("ascii")


def encrypted_ooxml() -> bytes:
    """The shape of a password-protected Office document: an OLE container with an EncryptedPackage stream."""
    header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504
    directory = "EncryptedPackage".encode("utf-16-le") + b"\x00" * 64 + "EncryptionInfo".encode("utf-16-le")
    return header + directory + b"\x00" * 1024


# --- Tables -------------------------------------------------------------------------------------------------------

PAGE_WIDTH = 595.0
PAGE_HEIGHT = 842.0


@dataclass(frozen=True)
class GridCell:
    """A cell of a drawn table: grid position, spans and Latin text (the standard PDF font has no Cyrillic)."""

    row: int
    col: int
    text: str
    row_span: int = 1
    col_span: int = 1


def _pdf_string(text: str) -> str:
    return "(" + text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ")"


def text_ops(x: float, top: float, text: str, size: float = 10) -> str:
    """Text at `x` whose line starts `top` points below the top of the page."""
    return f"BT /F1 {size} Tf {x:.2f} {PAGE_HEIGHT - top - size:.2f} Td {_pdf_string(text)} Tj ET\n"


def grid_ops(xs: list[float], top: float, row_height: float, cells: list[GridCell]) -> str:
    """A ruled table: every cell stroked as a rectangle in black, its text inside near the top-left corner."""
    ops = "0 G 0.5 w\n"
    for cell in cells:
        x0, x1 = xs[cell.col], xs[cell.col + cell.col_span]
        y_top = top + cell.row * row_height
        height = cell.row_span * row_height
        ops += f"{x0:.2f} {PAGE_HEIGHT - y_top - height:.2f} {x1 - x0:.2f} {height:.2f} re S\n"
        if cell.text:
            ops += text_ops(x0 + 3, y_top + 3, cell.text)
    return ops


def white_rect_ops(x0: float, top: float, x1: float, bottom: float) -> str:
    """A white filled rectangle, as word processors paint cell backgrounds: it draws nothing visible on white."""
    return f"1 g {x0:.2f} {PAGE_HEIGHT - bottom:.2f} {x1 - x0:.2f} {bottom - top:.2f} re f 0 g\n"


def rows_grid(rows: list[list[str]]) -> list[GridCell]:
    return [GridCell(r, c, text) for r, row in enumerate(rows) for c, text in enumerate(row)]


def pdf_pages(contents: list[str]) -> bytes:
    """A PDF whose pages are the given content streams, with Helvetica as /F1."""
    writer = PdfWriter()
    font = writer._add_object(
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
                NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
            }
        )
    )
    for ops in contents:
        page = writer.add_blank_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(ops.encode("latin-1"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def results_docx(rows: int, *, borders: bool = True) -> bytes:
    """A competition results table in DOCX: a two-row header whose «Задачи» cell spans two columns, the header rows
    repeated on every page, `rows` numbered result rows that do not break across pages. People are fictional."""
    document = docx.Document()
    document.add_paragraph("ИТОГОВЫЙ ПРОТОКОЛ")
    table = document.add_table(rows=2 + rows, cols=5)
    if borders:
        table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].merge(table.rows[1].cells[0]).text = "Место"
    header[1].merge(table.rows[1].cells[1]).text = "Участник"
    header[2].merge(header[3]).text = "Задачи"
    header[4].merge(table.rows[1].cells[4]).text = "Решено"
    table.rows[1].cells[2].text = "A"
    table.rows[1].cells[3].text = "B"
    for n in range(1, rows + 1):
        cells = table.rows[1 + n].cells
        for cell, value in zip(cells, [str(n), f"Участник Примерный {n}", "+", "+1", "2"], strict=True):
            cell.text = value
    for n, row in enumerate(table.rows):
        # Rows stay whole on a page; the first two are the header repeated on every page.
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if n < 2:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    document.add_paragraph("Главный судья соревнований")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def spans_docx() -> bytes:
    """A DOCX table with a horizontal merge, a vertical merge whose continuation holds text of its own, a row starting
    one column late and a nested table."""
    document = docx.Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Шапка на две колонки"
    table.cell(0, 2).merge(table.cell(1, 2)).text = "На две строки"
    # Word keeps text typed into the continuation of a vertical merge in its own element.
    continuation = table.rows[1]._tr.tc_lst[2]
    docx.table._Cell(continuation, table).paragraphs[0].add_run("и продолжение")
    table.cell(1, 0).text = "Левая"
    table.cell(1, 1).text = "Средняя"
    table.cell(2, 1).text = "После пропуска"
    table.cell(2, 2).add_table(rows=1, cols=2)
    nested = table.cell(2, 2).tables[0]
    nested.cell(0, 0).text = "Вложенная 1"
    nested.cell(0, 1).text = "Вложенная 2"
    # The last row starts at the second grid column: its first cell is removed and w:gridBefore takes its place.
    last = table.rows[2]._tr
    last.remove(last.tc_lst[0])
    grid_before = OxmlElement("w:gridBefore")
    grid_before.set(qn("w:val"), "1")
    last.get_or_add_trPr().insert(0, grid_before)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def docx_to_pdf(data: bytes, workdir: Path) -> bytes:
    """Converts a DOCX to PDF with LibreOffice, the way protocols are usually exported."""
    source = workdir / "table.docx"
    source.write_bytes(data)
    subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation={(workdir / 'profile').as_uri()}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(workdir),
            str(source),
        ],
        check=True,
        capture_output=True,
        env={"PATH": "/usr/bin:/bin", "HOME": str(workdir)},
    )
    return (workdir / "table.pdf").read_bytes()

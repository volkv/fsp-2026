"""Resume files of the template families: PDF (fpdf2, a Unicode TrueType font) and DOCX (python-docx).

Every family prints the same profile its own way; the labels do not depend on how (gold.py). Files are written with
fixed metadata dates, so a family renders a profile to the same bytes each time.
"""

import io
import zipfile
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path

import docx
from docx.shared import Pt
from fpdf import FPDF

from fsp_eval.families import FAMILY_BY_KEY, DateStyle
from fsp_eval.lexicon import LANGUAGE_NAMES, LEVEL_WORDS_HH
from fsp_eval.profiles import Job, Profile

FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")
FIXED_DATE = datetime(2026, 1, 1, tzinfo=UTC)

MONTHS_RU = (
    "Январь",
    "Февраль",
    "Март",
    "Апрель",
    "Май",
    "Июнь",
    "Июль",
    "Август",
    "Сентябрь",
    "Октябрь",
    "Ноябрь",
    "Декабрь",
)
MONTHS_EN = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
PRESENT: dict[DateStyle, str] = {
    "month_word_ru": "настоящее время",
    "numeric": "н. в.",
    "year": "по настоящее время",
    "month_word_en": "Present",
}


def format_date(style: DateStyle, value: tuple[int, int]) -> str:
    year, month = value
    match style:
        case "month_word_ru":
            return f"{MONTHS_RU[month - 1]} {year}"
        case "numeric":
            return f"{month:02d}.{year}"
        case "year":
            return str(year)
        case "month_word_en":
            return f"{MONTHS_EN[month - 1]} {year}"


def format_period(style: DateStyle, job: Job, dash: str) -> str:
    end = PRESENT[style] if job.end is None else format_date(style, job.end)
    return f"{format_date(style, job.start)} {dash} {end}"


def _plural(n: int, one: str, few: str, many: str) -> str:
    if n % 10 == 1 and n % 100 != 11:
        return one
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        return few
    return many


def duration_ru(months: int) -> str:
    years, rest = divmod(months, 12)
    parts = []
    if years:
        parts.append(f"{years} {_plural(years, 'год', 'года', 'лет')}")
    if rest:
        parts.append(f"{rest} {_plural(rest, 'месяц', 'месяца', 'месяцев')}")
    return " ".join(parts) or "меньше месяца"


def _job_months(job: Job, as_of: tuple[int, int]) -> int:
    end = job.end or as_of
    return (end[0] - job.start[0]) * 12 + end[1] - job.start[1] + 1


def language_lines(profile: Profile, style: str) -> list[str]:
    """`dash` «Английский — B2», `colon` «Английский язык: B2», `hh` «Английский — B2 — Средне-продвинутый»."""
    lines = []
    for code, level in profile.languages:
        ru, en = LANGUAGE_NAMES[code]
        match style:
            case "dash":
                lines.append(f"{ru} — {level}")
            case "colon":
                lines.append(f"{ru} язык: {level}")
            case "hh":
                lines.append(f"{ru} — {level} — {LEVEL_WORDS_HH[level]}")
            case "en":
                lines.append(f"{en} — {level}")
    lines.insert(0, "Russian — native" if style == "en" else "Русский — родной")
    return lines


def _dated(profile: Profile) -> bool:
    return profile.kind == "full"


# --- PDF ----------------------------------------------------------------------------------------------------------


class _Pdf(FPDF):
    def __init__(self) -> None:
        super().__init__(format="A4")
        self.add_font("DejaVu", "", str(FONT_DIR / "DejaVuSans.ttf"))
        self.add_font("DejaVu", "B", str(FONT_DIR / "DejaVuSans-Bold.ttf"))
        self.set_creation_date(FIXED_DATE)
        self.set_auto_page_break(auto=True, margin=15)
        self.add_page()

    def text_line(self, text: str, size: float = 10, bold: bool = False, gap: float = 0) -> None:
        self.set_font("DejaVu", "B" if bold else "", size)
        self.multi_cell(0, size * 0.5, text, new_x="LMARGIN", new_y="NEXT")
        if gap:
            self.ln(gap)

    def heading(self, text: str) -> None:
        self.ln(3)
        self.text_line(text, 13, bold=True, gap=1)

    def columns(self, left: float, right: float) -> None:
        self.set_left_margin(left)
        self.set_right_margin(right)
        self.set_x(left)


def _pdf_header(pdf: _Pdf, profile: Profile, contacts_inline: bool) -> None:
    pdf.text_line(profile.name, 18, bold=True, gap=1)
    pdf.text_line(profile.headline, 12, gap=1)
    if contacts_inline:
        pdf.text_line(f"{profile.city} · {profile.email} · {profile.phone}", 9, gap=2)


def _pdf_education(pdf: _Pdf, profile: Profile, heading: str) -> None:
    if not profile.education:
        return
    pdf.heading(heading)
    for entry in profile.education:
        pdf.text_line(entry.institution, 10, bold=True)
        pdf.text_line(f"{entry.degree}, {entry.year}", 9, gap=1)


def render_pdf_classic(profile: Profile, as_of: tuple[int, int]) -> bytes:
    del as_of
    pdf = _Pdf()
    _pdf_header(pdf, profile, contacts_inline=True)
    if profile.jobs:
        pdf.heading("Опыт работы")
        for job in profile.jobs:
            if _dated(profile):
                pdf.text_line(format_period("month_word_ru", job, "—"), 9)
            pdf.text_line(job.organization, 10)
            pdf.text_line(job.title, 11, bold=True)
            for line in job.lines:
                pdf.text_line(f"• {line.text}", 10)
            pdf.ln(2)
    pdf.heading("Ключевые навыки")
    pdf.text_line(", ".join(m.surface for m in profile.skills))
    _pdf_education(pdf, profile, "Образование")
    pdf.heading("Знание языков")
    for line in language_lines(profile, "dash"):
        pdf.text_line(line)
    pdf.heading("О себе")
    pdf.text_line(profile.about)
    return bytes(pdf.output())


def render_pdf_sidebar(profile: Profile, as_of: tuple[int, int]) -> bytes:
    del as_of
    pdf = _Pdf()
    _pdf_header(pdf, profile, contacts_inline=False)
    top = pdf.get_y() + 2
    # Left column: contacts, skills one per line, languages.
    pdf.columns(10, 145)
    pdf.set_y(top)
    pdf.set_x(10)
    pdf.heading("Контакты")
    for line in (profile.city, profile.email, profile.phone):
        pdf.text_line(line, 9)
    pdf.heading("Навыки")
    for mention in profile.skills:
        pdf.text_line(mention.surface, 9)
    pdf.heading("Языки")
    for line in language_lines(profile, "dash"):
        pdf.text_line(line, 9)
    # Right column: experience, education, about.
    pdf.columns(75, 10)
    pdf.set_y(top)
    pdf.set_x(75)
    if profile.jobs:
        pdf.heading("Опыт работы")
        for job in profile.jobs:
            pdf.text_line(job.title, 11, bold=True)
            pdf.text_line(job.organization, 10)
            if _dated(profile):
                pdf.text_line(format_period("year", job, "—"), 9)
            for line in job.lines:
                pdf.text_line(f"— {line.text}", 9)
            pdf.ln(2)
    _pdf_education(pdf, profile, "Образование")
    pdf.heading("Обо мне")
    pdf.text_line(profile.about, 9)
    return bytes(pdf.output())


def render_pdf_hh(profile: Profile, as_of: tuple[int, int]) -> bytes:
    pdf = _Pdf()
    pdf.text_line("Резюме обновлено 1 сентября 2026", 8, gap=2)
    _pdf_header(pdf, profile, contacts_inline=True)
    if profile.jobs:
        total = sum(_job_months(job, as_of) for job in profile.jobs)
        pdf.heading(f"Опыт работы — {duration_ru(total)}" if _dated(profile) else "Опыт работы")
        for job in profile.jobs:
            if pdf.get_y() > pdf.h - 70:
                pdf.add_page()
            top = pdf.get_y()
            pdf.columns(10, 150)
            pdf.set_xy(10, top)
            if _dated(profile):
                pdf.text_line(format_period("month_word_ru", job, "—").replace("настоящее", "по настоящее"), 8)
                pdf.text_line(duration_ru(_job_months(job, as_of)), 8)
            left_bottom = pdf.get_y()
            pdf.columns(65, 10)
            pdf.set_xy(65, top)
            pdf.text_line(job.organization, 10, bold=True)
            pdf.text_line(job.title, 10, bold=True, gap=1)
            for line in job.lines:
                pdf.text_line(line.text, 9)
            pdf.set_y(max(pdf.get_y(), left_bottom) + 3)
            pdf.columns(10, 10)
    pdf.columns(10, 10)
    pdf.heading("Навыки")
    pdf.text_line("   ".join(m.surface for m in profile.skills), 9)
    _pdf_education(pdf, profile, "Образование")
    pdf.heading("Знание языков")
    for line in language_lines(profile, "hh"):
        pdf.text_line(line, 9)
    pdf.heading("Дополнительная информация")
    pdf.text_line(profile.about, 9)
    return bytes(pdf.output())


# --- DOCX ---------------------------------------------------------------------------------------------------------


def _docx_document(profile: Profile) -> docx.document.Document:
    document = docx.Document()
    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.created = FIXED_DATE
    properties.modified = FIXED_DATE
    properties.revision = 1
    document.add_heading(profile.name, level=0)
    document.add_paragraph(profile.headline)
    return document


def _docx_bytes(document: docx.document.Document) -> bytes:
    """The package with every zip entry stamped with the fixed date: python-docx stamps the time of saving."""
    saved = io.BytesIO()
    document.save(saved)
    fixed = io.BytesIO()
    with zipfile.ZipFile(saved) as source, zipfile.ZipFile(fixed, "w") as target:
        for entry in source.infolist():
            stamped = zipfile.ZipInfo(entry.filename, date_time=FIXED_DATE.timetuple()[:6])
            stamped.compress_type = zipfile.ZIP_DEFLATED
            stamped.external_attr = entry.external_attr
            target.writestr(stamped, source.read(entry.filename))
    return fixed.getvalue()


def render_docx_headings(profile: Profile, as_of: tuple[int, int]) -> bytes:
    del as_of
    document = _docx_document(profile)
    document.add_paragraph(f"{profile.city}, {profile.phone}, {profile.email}")
    if profile.jobs:
        document.add_heading("Опыт работы", level=1)
        for job in profile.jobs:
            document.add_paragraph().add_run(job.title).bold = True
            document.add_paragraph(job.organization)
            if _dated(profile):
                document.add_paragraph(format_period("numeric", job, "–"))
            for line in job.lines:
                document.add_paragraph(line.text, style="List Bullet")
    document.add_heading("Навыки", level=1)
    for mention in profile.skills:
        document.add_paragraph(mention.surface, style="List Bullet")
    if profile.education:
        document.add_heading("Образование", level=1)
        for entry in profile.education:
            document.add_paragraph(entry.institution)
            document.add_paragraph(f"{entry.degree}, {entry.year}")
    document.add_heading("Иностранные языки", level=1)
    for line in language_lines(profile, "colon"):
        document.add_paragraph(line)
    document.add_heading("О себе", level=1)
    document.add_paragraph(profile.about)
    return _docx_bytes(document)


def render_docx_table(profile: Profile, as_of: tuple[int, int]) -> bytes:
    del as_of
    document = _docx_document(profile)
    document.add_paragraph(f"{profile.email} | {profile.phone}")
    if profile.jobs:
        document.add_heading("Work Experience", level=1)
        table = document.add_table(rows=0, cols=2)
        for job in profile.jobs:
            left, right = table.add_row().cells
            left.text = format_period("month_word_en", job, "–") if _dated(profile) else ""
            title = right.paragraphs[0]
            title.add_run(job.title).bold = True
            right.add_paragraph(job.organization)
            for line in job.lines:
                right.add_paragraph(f"• {line.text}").runs[0].font.size = Pt(10)
    document.add_heading("Skills", level=1)
    document.add_paragraph(", ".join(m.surface for m in profile.skills))
    if profile.education:
        document.add_heading("Education", level=1)
        for entry in profile.education:
            document.add_paragraph(entry.institution)
            document.add_paragraph(f"{entry.degree}, {entry.year}")
    document.add_heading("Languages", level=1)
    for line in language_lines(profile, "en"):
        document.add_paragraph(line)
    document.add_heading("About Me", level=1)
    document.add_paragraph(profile.about)
    return _docx_bytes(document)


RENDERERS: dict[str, Callable[[Profile, tuple[int, int]], bytes]] = {
    "pdf_classic": render_pdf_classic,
    "docx_headings": render_docx_headings,
    "pdf_sidebar": render_pdf_sidebar,
    "docx_table": render_docx_table,
    "pdf_hh": render_pdf_hh,
}


def render(profile: Profile, as_of: tuple[int, int]) -> tuple[bytes, str]:
    """The file of the profile in its family and the file extension."""
    return RENDERERS[profile.family](profile, as_of), FAMILY_BY_KEY[profile.family].file_format

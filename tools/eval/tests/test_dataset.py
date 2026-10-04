import dataclasses
import io
import json
import time
from collections import Counter
from pathlib import Path

import docx
import pytest

from fsp_eval.baselines import taxonomy_spellings
from fsp_eval.dataset import (
    DATASETS_DIR,
    DatasetError,
    build,
    labels_digest,
    load_manifest,
    new_manifest,
    write_documents,
)
from fsp_eval.families import FAMILIES, FAMILY_BY_KEY
from fsp_eval.gold import gold_of, period_key
from fsp_eval.lexicon import DISTRACTORS, ROLES, SKILL_SURFACES, fold_spelling
from fsp_eval.profiles import _has_cyrillic, assign, generate_profiles
from fsp_eval.render import format_period, render

MANIFEST = load_manifest("synthetic-resumes-v2")


@pytest.fixture(scope="module")
def dataset():
    return build(MANIFEST)


def test_the_committed_manifest_is_reproduced(dataset):
    assert len(dataset.profiles) == MANIFEST.profiles
    for split, facts in MANIFEST.splits.items():
        assert len(dataset.split(split)) == facts["profiles"]


def test_a_changed_generator_is_refused():
    tampered = dataclasses.replace(MANIFEST, labels_sha256="0" * 64)
    with pytest.raises(DatasetError, match="new dataset id"):
        build(tampered)


def test_generation_is_deterministic():
    first = generate_profiles(7, 40, (2026, 9), 30)
    assert first == generate_profiles(7, 40, (2026, 9), 30)
    assert labels_digest(first) != labels_digest(generate_profiles(8, 40, (2026, 9), 30))


def test_splits_share_no_profile_and_holdout_keeps_an_unseen_family(dataset):
    dev, holdout = dataset.split("dev"), dataset.split("holdout")
    assert {p.id for p in dev}.isdisjoint({p.id for p in holdout})
    # The English table layout is tuned on dev and measured on holdout profiles; `pdf_hh` is never seen while tuning.
    assert {p.family for p in dev} & {p.family for p in holdout} == {"docx_table"}
    assert "pdf_hh" in {p.family for p in holdout} - {p.family for p in dev}
    for profile in dataset.profiles:
        assert profile.split in FAMILY_BY_KEY[profile.family].splits
        assert assign(MANIFEST.seed, profile.id, MANIFEST.holdout_percent) == (
            profile.split,
            FAMILY_BY_KEY[profile.family],
        )
    assert {p.family for p in dataset.profiles} == {family.key for family in FAMILIES}


def test_the_dataset_covers_roles_grades_and_insufficient_resumes(dataset):
    kinds = Counter(p.kind for p in dataset.profiles)
    assert kinds["full"] > 200
    assert kinds["no_dates"] > 5
    assert kinds["skills_only"] > 5
    assert {p.role for p in dataset.profiles} == {role.key for role in ROLES}
    assert set(Counter(p.grade for p in dataset.profiles if p.kind == "full")) == {
        "intern",
        "junior",
        "middle",
        "senior",
        "lead",
    }


def test_every_skill_of_a_stack_can_be_written_in_english():
    for role in ROLES:
        for stack in role.stacks:
            for key in stack.skills:
                assert any(not _has_cyrillic(s) for s in SKILL_SURFACES[key]), key


def test_distractors_are_unknown_to_the_taxonomy():
    # «No skill» is the right answer for a distractor only while the taxonomy has no such spelling.
    known = taxonomy_spellings(Path(__file__).parent / "taxonomy")
    assert [surface for surface in DISTRACTORS if fold_spelling(surface) in known] == []


def test_labels_come_from_the_profile(dataset):
    for profile in dataset.profiles:
        gold = gold_of(profile)
        asserted = {m.key for m in profile.skills}
        # A skill the person only plans to learn is not a label.
        for mention in profile.negated:
            assert mention.key not in asserted
            assert mention.key not in gold.skills
        if profile.kind == "skills_only":
            assert (gold.role, gold.role_sufficient, gold.grade_sufficient) == (None, False, False)
            assert gold.periods == ()
        elif profile.kind == "no_dates":
            assert (gold.role, gold.role_sufficient, gold.grade_sufficient) == (profile.role, True, False)
            assert gold.periods == ()
        else:
            assert gold.grade == profile.grade
            assert len(gold.periods) == len(profile.jobs)
    typos = [m for p in dataset.profiles for m in p.skills if m.typo]
    assert typos
    assert all(m.surface not in SKILL_SURFACES[m.key] for m in typos)


def test_a_period_label_has_the_precision_the_template_prints():
    assert period_key((2021, 3), None, year_only=False) == "2021-03/present"
    assert period_key((2019, 1), (2021, 12), year_only=True) == "2019/2021"


def test_every_family_renders_the_profile(dataset, tmp_path):
    for family in FAMILIES:
        profile = next(p for p in dataset.profiles if p.family == family.key and p.kind == "full")
        content, extension = render(profile, MANIFEST.as_of_month)
        assert extension == family.file_format
        # Fixed metadata: the same profile gives the same file, whenever it is rendered. The second render comes
        # after a two-second boundary: zip entries store time to two seconds, PDF dates to one.
        time.sleep(2.1 - time.time() % 2)
        assert render(profile, MANIFEST.as_of_month)[0] == content
        if extension == "pdf":
            assert content.startswith(b"%PDF-")
            continue
        document = docx.Document(io.BytesIO(content))
        text = "\n".join(p.text for p in document.paragraphs)
        text += "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        for job in profile.jobs:
            assert job.organization in text
            assert job.title in text
            assert format_period(family.date_style, job, "–") in text
    files = write_documents(dataset, "holdout", tmp_path)
    assert len(files) == MANIFEST.splits["holdout"]["profiles"]
    assert all(f.path.is_file() and f.path.stat().st_size > 1000 for f in files)


def test_the_manifest_file_is_what_the_generator_prints():
    committed = json.loads((DATASETS_DIR / MANIFEST.id / "manifest.json").read_text(encoding="utf-8"))
    assert new_manifest(MANIFEST.id, MANIFEST.seed, MANIFEST.profiles, MANIFEST.as_of, MANIFEST.holdout_percent) == (
        committed
    )
